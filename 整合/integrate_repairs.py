#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1", "httpx>=0.25", "python-dotenv>=1.0"]
# ///
"""将 RTY 原始导出整合为每机型一个 sheet 的维修物料统计。

运行（在项目根目录，uv 会临时安装本脚本依赖，不修改项目依赖）：
    uv run 整合/integrate_repairs.py
    uv run 整合/integrate_repairs.py --concurrency 32 --overwrite
    uv run 整合/integrate_repairs.py --production-source api
    uv run 整合/integrate_repairs.py --quantity-field storage
    uv run 整合/integrate_repairs.py --offline --output /tmp/维修汇总.xlsx
    uv run 整合/integrate_repairs.py --compare-only

默认输入/输出为脚本所在目录的 原始数据.xlsx / 整合结果.xlsx。
保留原始工作表，在其后追加各机型统计页；统计页严格使用示例的12列。
生成后自动与 整合后示例.xlsx 对比，将结果写入 整合结果_对比报告.md。
示例只列出部分分组，按机型、维修维度、物料号定位比较，不要求行号相同。
示例有而原始数据没有的记录单独报告，不补造、不算作通过。
退出码：0=全部通过；1=运行失败；2=结果校验失败；3=仅存在示例与原始数据冲突。
网关地址依次取 --gateway-url、环境变量 DATA_GATEWAY_URL、backend/.env。
脚本独立通过 HTTP 读取本项目使用的数据源接口，不需要启动品质看板或连接 MongoDB。
默认同时查询16个生产订单，可通过 --concurrency 1..64 调整；每次运行输出总耗时。

统计口径：
* 分组键为 (机型, 维修维度3, 物料号)。物料按 ZCCLH、MATNR、
  RECORD01REPAIRM 顺序取第一个非空值；数字料号去掉 SAP 补齐的前导零。
* 一条维修记录计一次，不按 SN 去重物料次数；相同组输出 料号*n。
  不良现象、物料描述各自去重后换行展示。J 列为该组整机 SN 去重数。
* H 列为该机型在原始销售订单页全部生产订单的同料号 MENGE_A 之和，
  与本项目 LRR 的发料口径一致。它是发料量，不是 BOM 理论用量或维修耗用量。
  不用维修日期限制发料日期，也不把 MENGE_A 再乘整机数量。
* K 列默认汇总原始销售订单的“订单数量”(GAMNG)，可用
  --quantity-field storage 切换“入库数量”(WMENG)。--production-source api
  则从 /api/orders 查询同一批生产订单，保留相同数量口径。
* I=G/H，L=J/K，分母可用时采用与示例相同的 Excel 公式，
  打开 Excel/WPS 时自动重算，分母不可用时标记 n.a.。
  无物料、无发料记录、缺少数量或缺少 SN 时保留明确提示，不当成零。
* 示例文件中的 2400、1258 均为示例值；本脚本不硬编码这些数值。

接口失败或分页不完整会终止，不输出部分成功的统计。--offline 仅做本地
分组和整机数量统计，物料用量与部件故障率标记为未查询。
"""

from __future__ import annotations

import argparse
import asyncio
import math
import os
import re
import sys
import tempfile
from time import perf_counter
from copy import copy
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

try:
    import httpx
    from dotenv import dotenv_values
    from openpyxl import load_workbook
    from openpyxl.formula.translate import Translator
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties
except ImportError as exc:
    raise SystemExit(
        "缺少依赖，请使用 uv run 整合/integrate_repairs.py，"
        "或先安装 python -m pip install openpyxl httpx python-dotenv"
    ) from exc


SCRIPT_DIR = Path(__file__).resolve().parent
PAGE_SIZE = 100
MAX_API_ROWS = 250_000
HEADERS = (
    "序号",
    "机型",
    "不良现象维度3",
    "维修维度3",
    "物料号",
    "物料描述（短文本）",
    "部件故障计数",
    "物料使用量",
    "部件故障率",
    "整机故障数",
    "整机生产量",
    "整机故障率",
)
TEMPLATE_HEADERS = (
    "序号",
    "ZJXMC-机型",
    "ZNGWD3-不良现象维度3",
    "ZWXWD3-维修维度3",
    "ZCCLH-重插部件料号/MATNR-物料号/RECORD01REPAIRM-非关键件物料序号",
    "MAKTX-物料描述（短文本）",
    "ZCCLH-重插部件料号/MATNR-物料号/RECORD01REPAIRM-非关键件物料序号(计数)",
    "物料使用量",
    "部件故障率",
    "整机SN去重计数",
    "整机生产量",
    "整机故障率",
)
STAT_HEADER_VARIANTS = {HEADERS, TEMPLATE_HEADERS}
ALIASES = {
    "model": ("机型", "生产机型", "ZJXMC", "MAKTX_TH"),
    "order": ("生产订单", "生产订单号", "AUFNR"),
    "sales": ("销售订单", "VBELN"),
    "plant": ("下单工厂", "WERKS"),
    "order_qty": ("订单数量", "GAMNG"),
    "storage_qty": ("入库数量", "WMENG"),
    "repair": ("维修维度3", "ZWXWD3"),
    "symptom": ("不良现象维度3", "ZNGWD3"),
    "reinsert": ("重插部件料号", "ZCCLH"),
    "material": ("物料号", "MATNR"),
    "nonkey": ("非关键件物料序号", "RECORD01REPAIRM"),
    "description": ("物料描述（短文本）", "物料描述(短文本)", "MAKTX"),
    "sn": ("主机条码", "整机SN", "PCODE"),
}


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        if value.is_integer():
            return str(int(value))
    return str(value).strip()


def identifier(value: Any) -> str:
    value = text(value)
    if value.upper() in {"-", "--", "NULL", "NONE", "NAN", "N/A"}:
        return ""
    return value.lstrip("0") or "0" if value.isdigit() else value


def number(value: Any, label: str, *, signed: bool = False) -> Decimal | None:
    if text(value) == "":
        return None
    try:
        if isinstance(value, bool):
            raise InvalidOperation
        result = Decimal(text(value).replace(",", ""))
        if not result.is_finite() or (not signed and result < 0):
            raise InvalidOperation
        return result
    except InvalidOperation as exc:
        raise ValueError(f"{label}不是有效数量：{value!r}") from exc


def excel_number(value: Decimal | None) -> int | float | None:
    if value is None:
        return None
    return int(value) if value == value.to_integral_value() else float(value)


def source_data(item: dict[str, Any]) -> dict[str, Any]:
    raw = item.get("raw")
    if not isinstance(raw, dict):
        return item
    return raw["data"] if isinstance(raw.get("data"), dict) else raw


def read_table(
    workbook: Any, required: set[str]
) -> tuple[str, list[tuple[int, dict[str, Any]]]]:
    candidates = []
    for sheet in workbook:
        for row in sheet.iter_rows(max_row=min(20, sheet.max_row)):
            columns = {}
            for index, cell in enumerate(row):
                heading = text(cell.value)
                for key, aliases in ALIASES.items():
                    if heading in aliases or heading.split("-", 1)[0] in aliases:
                        if key in columns:
                            raise ValueError(f"{sheet.title} 表头重复：{heading}")
                        columns[key] = index
            if required <= columns.keys():
                records = [
                    (row_number, {key: values[index] for key, index in columns.items()})
                    for row_number, values in enumerate(
                        sheet.iter_rows(min_row=row[0].row + 1, values_only=True),
                        start=row[0].row + 1,
                    )
                    if any(text(value) for value in values)
                ]
                candidates.append((sheet.title, records))
                break
    if len(candidates) != 1:
        raise ValueError(
            f"应恰好找到一个包含 {', '.join(sorted(required))} 表头的工作表，"
            f"实际找到 {len(candidates)} 个"
        )
    return candidates[0]


@dataclass
class Order:
    number: str
    model: str
    order_qty: Decimal | None
    storage_qty: Decimal | None
    sales: str
    plant: str


@dataclass
class RepairGroup:
    count: int = 0
    symptoms: dict[str, None] = field(default_factory=dict)
    descriptions: dict[str, None] = field(default_factory=dict)
    sns: set[str] = field(default_factory=set)
    missing_sn: int = 0
    source_rows: list[int] = field(default_factory=list)


@dataclass
class Source:
    models: list[str] = field(default_factory=list)
    orders: dict[str, Order] = field(default_factory=dict)
    groups: dict[tuple[str, str, str], RepairGroup] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    repair_count: int = 0
    repair_sheet: str = ""


def load_source(path: Path) -> Source:
    result = Source()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        order_sheet, orders = read_table(workbook, {"order", "model", "order_qty"})
        for row_number, row in orders:
            label = f"{order_sheet} 第 {row_number} 行"
            order_number, model = identifier(row["order"]), text(row["model"])
            if not order_number or not model:
                raise ValueError(f"{label}缺少生产订单或机型")
            order = Order(
                order_number,
                model,
                number(row["order_qty"], f"{label}订单数量"),
                number(row.get("storage_qty"), f"{label}入库数量"),
                identifier(row.get("sales")),
                text(row.get("plant")),
            )
            if order_number in result.orders:
                if result.orders[order_number] != order:
                    raise ValueError(
                        f"{label}生产订单 {order_number} 重复且内容不同，无法确定数量"
                    )
                result.warnings.append(f"{label}重复生产订单 {order_number}，仅计一次")
                continue
            result.orders[order_number] = order
            if model not in result.models:
                result.models.append(model)
        if not result.orders:
            raise ValueError("销售订单页没有有效生产订单")

        repair_sheet, repairs = read_table(workbook, {"model", "repair"})
        result.repair_sheet = repair_sheet
        if repairs and not {"reinsert", "material", "nonkey"} & repairs[0][1].keys():
            raise ValueError(f"{repair_sheet}缺少三个物料号表头")
        for row_number, row in repairs:
            label = f"{repair_sheet} 第 {row_number} 行"
            model = text(row["model"])
            if not model:
                raise ValueError(f"{label}缺少机型")
            order_number = identifier(row.get("order"))
            order = result.orders.get(order_number)
            if order is None or order.model != model:
                raise ValueError(
                    f"{label}的生产订单与销售订单页无法精确关联：{order_number}/{model}"
                )
            materials = [
                value
                for key in ("reinsert", "material", "nonkey")
                if (value := identifier(row.get(key)))
            ]
            if len(set(materials)) > 1:
                result.warnings.append(
                    f"{label}多个物料字段有值，按优先级选用 {materials[0]}"
                )
            material = materials[0] if materials else ""
            repair = text(row["repair"])
            if not material:
                result.warnings.append(f"{label}未填物料号，保留记录但不计算部件故障率")
            if not repair:
                result.warnings.append(f"{label}未填维修维度3，保留为单独分组")
            group = result.groups.setdefault((model, repair, material), RepairGroup())
            group.count += 1
            group.source_rows.append(row_number)
            if symptom := text(row.get("symptom")):
                group.symptoms[symptom] = None
            # MAKTX describes MATNR. Do not attach it to a different selected PN.
            described_material = identifier(row.get("material"))
            if (
                (description := text(row.get("description")))
                and material
                and (not described_material or described_material == material)
            ):
                group.descriptions[description] = None
            if sn := text(row.get("sn")):
                group.sns.add(sn)
            else:
                group.missing_sn += 1
                result.warnings.append(f"{label}未填整机SN，该分组不计算整机故障率")
            result.repair_count += 1
    finally:
        workbook.close()
    return result


class Gateway:
    """Bounded, complete pagination over the same source APIs used by the BFF."""

    def __init__(self, client: httpx.AsyncClient):
        self.client = client

    async def rows(self, path: str, **filters: str) -> list[dict[str, Any]]:
        result: list[dict[str, Any]] = []
        expected_total = None
        page = 1
        while True:
            response = await self.client.get(
                path, params={**filters, "page": page, "pageSize": PAGE_SIZE}
            )
            response.raise_for_status()
            payload = response.json()
            if not isinstance(payload, dict):
                raise ValueError(f"{path} 返回格式错误")
            items, total = payload.get("items"), payload.get("total")
            if items is None and total == 0:
                items = []
            if (
                type(total) is not int
                or not 0 <= total <= MAX_API_ROWS
                or not isinstance(items, list)
                or any(not isinstance(item, dict) for item in items)
                or len(items) != min(PAGE_SIZE, max(0, total - len(result)))
            ):
                raise ValueError(f"{path} 分页数据不完整或超过 {MAX_API_ROWS} 行")
            if expected_total is not None and expected_total != total:
                raise ValueError(f"{path} 分页期间总数变化，请重试")
            expected_total = total
            result.extend(items)
            if len(result) == total:
                return result
            page += 1


@dataclass
class Materials:
    quantities: dict[tuple[str, str], Decimal] = field(default_factory=dict)
    descriptions: dict[tuple[str, str], dict[str, None]] = field(default_factory=dict)
    incomplete_models: set[str] = field(default_factory=set)
    row_count: int = 0


def aggregate_bom(
    source: Source, batches: list[tuple[str, list[dict[str, Any]]]]
) -> Materials:
    result = Materials()
    seen: dict[str, tuple[Any, ...]] = {}
    for order_number, items in batches:
        order = source.orders[order_number]
        matched = 0
        for item in items:
            row = source_data(item)
            if identifier(row.get("AUFNR_1")) != order_number:
                continue
            if order.sales and identifier(row.get("VBELN_EX")) != order.sales:
                raise ValueError(f"生产订单 {order_number} 的发料记录销售订单不一致")
            material = identifier(row.get("MATNR"))
            quantity = number(
                row.get("MENGE_A"),
                f"订单 {order_number} 物料 {material} MENGE_A",
                signed=True,
            )
            if not material or quantity is None:
                raise ValueError(f"生产订单 {order_number} 的发料记录缺少料号或数量")
            identity = text(item.get("id") or row.get("_source_key") or row.get("_id"))
            if not identity:
                raise ValueError(
                    f"生产订单 {order_number} 的发料记录缺少唯一ID，无法防止重复计数"
                )
            signature = (order_number, material, quantity, text(row.get("MAKTX")))
            matched += 1
            if identity in seen:
                if seen[identity] != signature:
                    raise ValueError(f"发料记录 {identity} 重复且内容不同")
                continue
            seen[identity] = signature
            key = (order.model, material)
            result.quantities[key] = result.quantities.get(key, Decimal(0)) + quantity
            if description := text(row.get("MAKTX")):
                result.descriptions.setdefault(key, {})[description] = None
            result.row_count += 1
        if not matched:
            result.incomplete_models.add(order.model)
            source.warnings.append(
                f"生产订单 {order_number} 未查到发料记录，{order.model} 用量标为不完整"
            )
    return result


async def fetch_sources(source: Source, args: argparse.Namespace) -> Materials:
    config = dotenv_values(args.env_file) if args.env_file.is_file() else {}
    gateway_url = (
        args.gateway_url
        or os.getenv("DATA_GATEWAY_URL")
        or config.get("DATA_GATEWAY_URL")
        or ""
    ).rstrip("/")
    if not gateway_url:
        raise ValueError(
            "未配置 DATA_GATEWAY_URL，请使用 --gateway-url 或配置 backend/.env"
        )
    concurrency = min(args.concurrency, len(source.orders))
    semaphore = asyncio.Semaphore(concurrency)
    print(f"数据源查询并发数：{concurrency}", flush=True)
    async with httpx.AsyncClient(
        base_url=gateway_url,
        timeout=httpx.Timeout(args.timeout, connect=min(args.timeout, 10)),
        trust_env=False,
        limits=httpx.Limits(
            max_connections=concurrency, max_keepalive_connections=concurrency
        ),
    ) as client:
        gateway = Gateway(client)
        completed = 0

        async def fetch(order: Order) -> tuple[str, list[dict[str, Any]]]:
            nonlocal completed
            async with semaphore:
                # Older gateways ignore the productionOrder view filter.
                # Keyword recalls candidates; aggregate_bom enforces exact ownership.
                rows = await gateway.rows(
                    "/api/views/ZSGV_ZSD124", keyword=order.number
                )
                if args.production_source == "api":
                    candidates = await gateway.rows("/api/orders", keyword=order.number)
                    matches = [
                        item
                        for item in candidates
                        if identifier(
                            item.get("aufnr") or source_data(item).get("AUFNR")
                        )
                        == order.number
                    ]
                    if len(matches) != 1:
                        raise ValueError(
                            f"订单 {order.number} 在接口匹配到 {len(matches)} 条，无法确定数量"
                        )
                    item = matches[0]
                    raw = source_data(item)
                    if (
                        text(item.get("productionModel") or raw.get("MAKTX_TH"))
                        != order.model
                    ):
                        raise ValueError(
                            f"订单 {order.number} 接口机型与原始文件不一致"
                        )
                    if (
                        order.sales
                        and identifier(item.get("salesOrder") or raw.get("VBELN"))
                        != order.sales
                    ):
                        raise ValueError(
                            f"订单 {order.number} 接口销售订单与原始文件不一致"
                        )
                    order.order_qty = number(
                        item.get("orderQuantity", raw.get("GAMNG")),
                        f"订单 {order.number}订单数量",
                    )
                    order.storage_qty = number(
                        item.get("storageQuantity", raw.get("WMENG")),
                        f"订单 {order.number}入库数量",
                    )
                completed += 1
                if completed % 10 == 0 or completed == len(source.orders):
                    print(
                        f"已查询生产订单：{completed}/{len(source.orders)}", flush=True
                    )
                return order.number, rows

        tasks = [asyncio.create_task(fetch(order)) for order in source.orders.values()]
        try:
            batches = await asyncio.gather(*tasks)
        finally:
            for task in tasks:
                if not task.done():
                    task.cancel()
            await asyncio.gather(*tasks, return_exceptions=True)
    return aggregate_bom(source, batches)


def total_production(source: Source, model: str, quantity_field: str) -> Decimal | None:
    values = [
        order.order_qty if quantity_field == "order" else order.storage_qty
        for order in source.orders.values()
        if order.model == model
    ]
    return (
        sum(values, Decimal(0))
        if values and all(v is not None for v in values)
        else None
    )


def material_usage(
    materials: Materials | None, model: str, material: str
) -> int | float | str:
    if not material:
        return "未填写物料号"
    if materials is None:
        return "未查询（离线）"
    if model in materials.incomplete_models:
        return "发料记录不完整"
    quantity = materials.quantities.get((model, material))
    return "未匹配发料记录" if quantity is None else excel_number(quantity)


def set_text(cell: Any, value: str) -> None:
    if len(value) > 32_767:
        raise ValueError(
            f"{cell.coordinate} 文本超过 Excel 单元格限制，请缩小原始数据范围"
        )
    cell.value = value
    cell.data_type = "s"


def rate_formula(
    row: int, numerator: str, denominator: str, count: Any, quantity: Any
) -> str:
    if type(count) in {int, float} and type(quantity) in {int, float} and quantity > 0:
        return f"={numerator}{row}/{denominator}{row}"
    return "n.a."


def copy_cell_style(target: Any, reference: Any) -> None:
    # Style IDs belong to their workbook; register the actual style components.
    for attribute in (
        "font",
        "fill",
        "border",
        "alignment",
        "protection",
        "number_format",
    ):
        setattr(target, attribute, copy(getattr(reference, attribute)))


def clear_fill(cell: Any) -> None:
    cell.fill = PatternFill(fill_type=None)


def write_output(
    source: Source, materials: Materials | None, args: argparse.Namespace
) -> None:
    workbook = load_workbook(args.input)
    template = load_workbook(args.example)
    template_sheets = [
        sheet
        for sheet in template
        if tuple(text(cell.value) for cell in sheet[1][:12]) in STAT_HEADER_VARIANTS
    ]
    reference = template_sheets[0] if template_sheets else None
    model_order = list(
        dict.fromkeys(
            [sheet.title for sheet in template_sheets if sheet.title in source.models]
            + source.models
        )
    )
    workbook.calculation = CalcProperties(
        calcId=0, fullCalcOnLoad=True, forceFullCalc=True
    )
    workbook.properties.description = f"原始数据：{args.input.name}；物料使用量：ZSGV_ZSD124.MENGE_A，按原始生产订单关联"
    used_names = {name.casefold() for name in workbook.sheetnames}
    widths = (7, 17, 76, 22, 38, 46, 24, 22, 15, 20, 18, 15)
    for model in model_order:
        sheet_name = re.sub(r"[\\/*?:\[\]]", "_", model)[:31]
        base_name = sheet_name
        suffix = 2
        while sheet_name.casefold() in used_names:
            sheet_name = f"{base_name[:27]}_{suffix}"
            suffix += 1
        used_names.add(sheet_name.casefold())
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(list(HEADERS))
        sheet.freeze_panes = "C2"
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        for cell in sheet[1]:
            if reference:
                copy_cell_style(cell, reference.cell(1, cell.column))
            else:
                cell.font = Font(name="宋体", size=11, bold=True)
            clear_fill(cell)
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        sheet.row_dimensions[1].height = 105
        groups = sorted(
            (
                (repair, material, group)
                for (m, repair, material), group in source.groups.items()
                if m == model
            ),
            key=lambda item: (-item[2].count, item[0], item[1]),
        )
        production = total_production(source, model, args.quantity_field)
        for index, (repair, material, group) in enumerate(groups, 1):
            row_number = index + 1
            descriptions = dict(group.descriptions)
            if materials is not None:
                descriptions.update(materials.descriptions.get((model, material), {}))
            usage = material_usage(materials, model, material)
            sn_count = (
                f"缺少SN（{group.missing_sn}条）"
                if group.missing_sn
                else len(group.sns)
            )
            production_value = (
                excel_number(production) if production is not None else "缺少整机数量"
            )
            values = [
                index,
                model,
                "\n".join(group.symptoms),
                repair or "未填写维修维度3",
                f"{material or '未填写物料号'}*{group.count}",
                "\n".join(descriptions),
                group.count,
                usage,
                rate_formula(row_number, "G", "H", group.count, usage),
                sn_count,
                production_value,
                rate_formula(row_number, "J", "K", sn_count, production_value),
            ]
            sheet.append(values)
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row_number, column)
                if isinstance(value, str) and not (
                    column in {9, 12} and value.startswith("=")
                ):
                    set_text(cell, value)
                if reference:
                    copy_cell_style(cell, reference.cell(2, column))
                else:
                    cell.font = Font(name="宋体", size=12)
                clear_fill(cell)
                cell.alignment = Alignment(
                    horizontal="right"
                    if column in {1, 7, 8, 9, 10, 11, 12}
                    else "left",
                    vertical="top",
                    wrap_text=True,
                )
                if column in {9, 12}:
                    cell.number_format = "0.00%"
                elif column in {1, 7, 8, 10, 11}:
                    cell.number_format = (
                        "#,##0.00######" if isinstance(value, float) else "#,##0"
                    )
            line_count = max(
                sum(
                    max(1, math.ceil(len(line) / limit))
                    for line in text(value).split("\n")
                )
                for value, limit in ((values[2], 33), (values[3], 10), (values[5], 20))
            )
            sheet.row_dimensions[row_number].height = max(30, line_count * 18 + 6)
        if not groups:
            for coordinate, value in {
                "B2": model,
                "D2": "无维修故障记录",
                "G2": 0,
                "J2": 0,
                "K2": excel_number(production)
                if production is not None
                else "缺少整机数量",
                "L2": rate_formula(2, "J", "K", 0, excel_number(production)),
            }.items():
                sheet[coordinate] = value
            set_text(sheet["B2"], model)
            for cell in sheet[2][:12]:
                cell.font = Font(name="宋体", size=11)
                cell.alignment = Alignment(
                    horizontal="right" if cell.column >= 7 else "left", vertical="top"
                )
                if cell.column in {7, 10, 11}:
                    cell.number_format = "#,##0"
            sheet["L2"].number_format = "0.00%"
            sheet.row_dimensions[2].height = 30
        sheet.auto_filter.ref = f"A1:L{max(2, len(groups) + 1)}"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"
        sheet.print_options.horizontalCentered = True
        sheet.print_area = sheet.auto_filter.ref

    template.close()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=args.output.parent, suffix=".xlsx", delete=False
        ) as handle:
            temporary_path = Path(handle.name)
        workbook.save(temporary_path)
        temporary_path.replace(args.output)
    finally:
        workbook.close()
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


@dataclass
class Comparison:
    checks: list[tuple[str, bool, str]] = field(default_factory=list)
    conflicts: list[str] = field(default_factory=list)
    samples: list[tuple[str, str, str, bool]] = field(default_factory=list)
    unavailable: list[str] = field(default_factory=list)

    @property
    def failures(self) -> int:
        return sum(not passed for _, passed, _ in self.checks)

    @property
    def exit_code(self) -> int:
        return 2 if self.failures else 3 if self.conflicts else 0

    def check(self, name: str, passed: bool, detail: str = "") -> None:
        self.checks.append((name, passed, detail))


def material_count(value: Any) -> tuple[str, int]:
    parts = text(value).rsplit("*", 1)
    if len(parts) != 2 or not parts[1].isdigit() or int(parts[1]) < 1:
        raise ValueError(f"物料号计数格式错误：{value!r}，应为 料号*n")
    material = "" if parts[0] == "未填写物料号" else identifier(parts[0])
    return material, int(parts[1])


def cell_value(value: Any) -> Any:
    return text(value) if isinstance(value, str) or value is None else value


def compare_output(
    source: Source, args: argparse.Namespace, materials: Materials | None = None
) -> Comparison:
    result = Comparison()
    output = load_workbook(args.output)
    original = load_workbook(args.input)
    example = load_workbook(args.example)
    try:
        for sheet in original:
            present = sheet.title in output.sheetnames
            differences = []
            if present:
                target = output[sheet.title]
                if (target.max_row, target.max_column) != (
                    sheet.max_row,
                    sheet.max_column,
                ):
                    differences.append("工作表范围改变")
                for row in sheet:
                    for cell in row:
                        actual = target[cell.coordinate]
                        if cell_value(actual.value) != cell_value(cell.value):
                            differences.append(f"{cell.coordinate} 原始值改变")
            result.check(
                f"保留原始页 {sheet.title}",
                present and not differences,
                "原始值与公式均保留；不检查样式"
                if present and not differences
                else "；".join(differences[:5]) or "缺少原始页",
            )
        result.check(
            "原始页顺序",
            output.sheetnames[: len(original.sheetnames)] == original.sheetnames,
            "原始页位于各机型统计页之前，与示例结构一致",
        )

        generated = [
            sheet for sheet in output if sheet.title not in original.sheetnames
        ]
        by_model = {}
        for sheet in generated:
            model = text(sheet["B2"].value)
            result.check(
                f"{sheet.title} 机型唯一",
                model not in by_model and model in source.models,
            )
            by_model[model] = sheet
        result.check(
            "机型覆盖",
            set(by_model) == set(source.models),
            f"应有 {len(source.models)} 个机型页",
        )

        actual_groups = {}
        counted = 0
        for model, sheet in by_model.items():
            result.check(
                f"{model} 12列表头",
                sheet.max_column == 12
                and tuple(text(cell.value) for cell in sheet[1]) == HEADERS,
                "使用中文说明列名，顺序一致且无额外说明列",
            )
            expected = {key for key in source.groups if key[0] == model}
            failures = []
            for row in sheet.iter_rows(min_row=2, max_col=12):
                if not expected and row[3].value == "无维修故障记录":
                    continue
                try:
                    material, occurrences = material_count(row[4].value)
                except ValueError as exc:
                    failures.append(str(exc))
                    continue
                repair = text(row[3].value)
                key = (model, "" if repair == "未填写维修维度3" else repair, material)
                if key in actual_groups or key not in expected:
                    failures.append(
                        f"第{row[0].row}行分组重复或原始记录不存在：{key[1:]}"
                    )
                    continue
                actual_groups[key] = row
                group = source.groups[key]
                counted += row[6].value if type(row[6].value) is int else 0
                checks = [
                    (occurrences == group.count == row[6].value, "物料计数"),
                    (text(row[1].value) == model, "机型"),
                    (
                        set(text(row[2].value).splitlines()) == set(group.symptoms),
                        "不良现象",
                    ),
                    (row[0].value == row[0].row - 1, "序号"),
                    (
                        row[9].value
                        == (
                            f"缺少SN（{group.missing_sn}条）"
                            if group.missing_sn
                            else len(group.sns)
                        ),
                        "SN去重",
                    ),
                    (
                        row[10].value
                        == (
                            excel_number(
                                total_production(source, model, args.quantity_field)
                            )
                            if total_production(source, model, args.quantity_field)
                            is not None
                            else "缺少整机数量"
                        ),
                        "整机数量",
                    ),
                    (
                        row[8].value
                        == rate_formula(
                            row[0].row, "G", "H", row[6].value, row[7].value
                        ),
                        "部件故障率公式",
                    ),
                    (
                        row[11].value
                        == rate_formula(
                            row[0].row, "J", "K", row[9].value, row[10].value
                        ),
                        "整机故障率公式",
                    ),
                ]
                if materials is not None:
                    checks.append(
                        (
                            row[7].value == material_usage(materials, model, material),
                            "接口发料量",
                        )
                    )
                for passed, label in checks:
                    if not passed:
                        failures.append(f"第{row[0].row}行{label}不一致")
                if not isinstance(row[7].value, (int, float)):
                    result.unavailable.append(
                        f"{model} / {key[1]} / {material or '未填写物料号'}：{row[7].value}"
                    )
            keys = {key for key in actual_groups if key[0] == model}
            result.check(
                f"{model} 全部分组覆盖",
                keys == expected,
                f"{len(keys)}/{len(expected)} 个分组",
            )
            result.check(
                f"{model} 原始数据及计算核对",
                not failures,
                "计数、SN去重、整机数量与故障率公式一致"
                if not failures
                else "；".join(failures[:10]),
            )
            if not expected:
                row = sheet[2]
                result.check(
                    f"{model} 无故障机型",
                    row[6].value == row[9].value == 0
                    and row[10].value
                    == excel_number(
                        total_production(source, model, args.quantity_field)
                    )
                    and row[11].value == rate_formula(2, "J", "K", 0, row[10].value),
                    "保留机型，故障计数为0，整机数量来自其订单",
                )
        result.check(
            "维修记录总数对账",
            counted == source.repair_count,
            f"结果计数 {counted} / 原始记录 {source.repair_count}",
        )

        reference_sheets = [
            sheet
            for sheet in example
            if tuple(text(cell.value) for cell in sheet[1][:12]) in STAT_HEADER_VARIANTS
        ]
        result.check("找到示例统计页", bool(reference_sheets), "按12列表头识别机型样例")
        for sheet in reference_sheets:
            model = sheet.title
            for row in sheet.iter_rows(min_row=2, max_col=12):
                if not text(row[4].value):
                    continue
                model = text(row[1].value) or model
                material, occurrences = material_count(row[4].value)
                key = (model, text(row[3].value), material)
                label = f"{sheet.title}!E{row[0].row} {key[1]} / {row[4].value}"
                if key not in source.groups:
                    result.conflicts.append(
                        f"{label}：在原始“{source.repair_sheet}”中该机型、维修维度、料号组合为0条，"
                        f"示例要求{occurrences}条。没有补造此分组。"
                    )
                    continue
                actual = actual_groups.get(key)
                if actual is None:
                    result.check(label, False, "原始数据有此分组，但结果缺失")
                    continue
                source_group = source.groups[key]
                result.check(
                    f"{label} 来源计数",
                    occurrences == source_group.count,
                    f"示例{occurrences}条；原始{source_group.count}条，原始行号 {source_group.source_rows}",
                )
                for expected_cell, actual_cell in zip(row, actual, strict=True):
                    value = expected_cell.value
                    if value is None:
                        continue
                    expected_value = value
                    if expected_cell.data_type == "f":
                        expected_value = Translator(
                            value, origin=expected_cell.coordinate
                        ).translate_formula(actual_cell.coordinate)
                    if expected_cell.column == 3:
                        passed = set(text(expected_value).splitlines()) == set(
                            text(actual_cell.value).splitlines()
                        )
                    else:
                        passed = cell_value(expected_value) == cell_value(
                            actual_cell.value
                        )
                    position = f"{sheet.title}!{expected_cell.coordinate} → {actual_cell.parent.title}!{actual_cell.coordinate}"
                    result.samples.append(
                        (
                            position,
                            text(expected_value),
                            text(actual_cell.value),
                            passed,
                        )
                    )
                    result.check(position, passed, "按分组定位；公式平移后比较")
    finally:
        output.close()
        original.close()
        example.close()
    return result


def write_comparison_report(
    result: Comparison, source: Source, args: argparse.Namespace
) -> None:
    def escape(value: str) -> str:
        return value.replace("|", "\\|").replace("\n", "<br>")

    status = (
        "结果校验失败"
        if result.failures
        else "可验证项通过，示例与原始数据存在冲突，不能认定完全一致"
        if result.conflicts
        else "全部对比通过"
    )
    lines = [
        "# 整合结果与示例对比",
        "",
        f"结论：**{status}**。",
        "",
        f"输入：{args.input.name}；示例：{args.example.name}；输出：{args.output.name}。",
        f"原始 {source.repair_count} 条维修记录、{len(source.orders)} 个生产订单、{len(source.models)} 个机型。",
        f"检查 {len(result.checks)} 项，通过 {len(result.checks) - result.failures} 项，失败 {result.failures} 项，"
        f"示例与原始数据冲突 {len(result.conflicts)} 项。",
        "",
        "## 示例与原始数据差异",
        "",
        *([f"- {item}" for item in result.conflicts] or ["无。"]),
        "",
        "示例仅展示部分分组。按“机型 + 维修维度3 + 物料号”匹配，不以行号或总行数判断相同，"
        "示例空白单元格不作为要求结果为空的依据。示例有、原始数据没有的分组不补造，也不计为通过。",
        "",
        "## 示例单元格逐项对比",
        "",
        "| 示例位置 → 结果位置 | 预期值或平移后公式 | 实际值 | 结果 |",
        "| --- | --- | --- | --- |",
        *[
            f"| {escape(position)} | {escape(expected)} | {escape(actual)} | {'通过' if passed else '失败'} |"
            for position, expected, actual, passed in result.samples
        ],
        "",
        "## 完整性与计算检查",
        "",
        "| 检查项 | 结果 | 说明 |",
        "| --- | --- | --- |",
        *[
            f"| {escape(name)} | {'通过' if passed else '失败'} | {escape(detail)} |"
            for name, passed, detail in result.checks
            if " → " not in name
        ],
        "",
        "## 数据缺项",
        "",
        *([f"- {escape(item)}" for item in result.unavailable] or ["无。"]),
        "",
        "## 统计口径",
        "",
        "- 物料字段按 ZCCLH、MATNR、RECORD01REPAIRM 优先级取第一个非空值，数字料号移除SAP补零。",
        "- 物料使用量按原始销售订单中的生产订单精确关联 MENGE_A 发料量，不再乘整机数量。",
        f"- 整机分母：{'订单数量/GAMNG' if args.quantity_field == 'order' else '入库数量/WMENG'}；"
        f"来源：{'接口' if args.production_source == 'api' else '原始订单页'}。",
        "- 部件故障率=物料计数/物料使用量；整机故障率=该分组SN去重数/整机数量。",
        "- 各组SN可能交叉，不应将各行SN计数相加当作机型总故障机器数。",
        "- 分母不可用时标记 n.a.。公式及引用已检查，实际重算由打开文件的Excel/WPS完成。",
        "- 本次验收只检查数据、结构和计算，不检查字体、颜色、列宽、行高等样式。",
        "- --compare-only 不重新请求接口；接口用量的全部逐组检查在生成流程中完成。",
        "",
    ]
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text("\n".join(lines), encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--input", type=Path, default=SCRIPT_DIR / "原始数据.xlsx")
    parser.add_argument("--output", type=Path, default=SCRIPT_DIR / "整合结果.xlsx")
    parser.add_argument("--example", type=Path, default=SCRIPT_DIR / "整合后示例.xlsx")
    parser.add_argument("--report", type=Path, help="默认写入 输出文件名_对比报告.md")
    parser.add_argument(
        "--compare-only", action="store_true", help="仅核对已有结果，不请求数据源"
    )
    parser.add_argument(
        "--env-file", type=Path, default=SCRIPT_DIR.parent / "backend" / ".env"
    )
    parser.add_argument(
        "--gateway-url", help="本项目数据源网关地址（不是品质看板4001端口）"
    )
    parser.add_argument(
        "--production-source", choices=("excel", "api"), default="excel"
    )
    parser.add_argument(
        "--quantity-field", choices=("order", "storage"), default="order"
    )
    parser.add_argument(
        "--offline", action="store_true", help="不请求接口，仅做原始文件分组"
    )
    parser.add_argument("--timeout", type=float, default=120)
    parser.add_argument(
        "--concurrency",
        type=int,
        choices=range(1, 65),
        default=16,
        metavar="1..64",
        help="同时查询的生产订单数（默认16，连接池容量同步调整）",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="允许替换已有结果（始终保护输入和示例文件）",
    )
    args = parser.parse_args()
    args.input, args.output = args.input.resolve(), args.output.resolve()
    args.example = args.example.resolve()
    args.report = (
        args.report or args.output.with_name(f"{args.output.stem}_对比报告.md")
    ).resolve()
    if args.timeout <= 0 or not math.isfinite(args.timeout):
        parser.error("--timeout 必须是大于0的有限数")
    if args.offline and args.production_source == "api":
        parser.error("--offline 不能与 --production-source api 同时使用")
    if not args.input.is_file():
        parser.error(f"输入文件不存在：{args.input}")
    if not args.example.is_file():
        parser.error(f"示例文件不存在：{args.example}")
    if args.output == args.input or args.output == args.example:
        parser.error("输出不能覆盖原始输入文件或整合后示例.xlsx")
    if args.report.suffix.lower() != ".md" or args.report in {
        args.input,
        args.output,
        args.example,
    }:
        parser.error("对比报告必须是独立的 .md 文件")
    if args.output.suffix.lower() != ".xlsx":
        parser.error("输出文件扩展名必须为 .xlsx")
    if args.compare_only and not args.output.is_file():
        parser.error(f"待核对结果不存在：{args.output}")
    if args.output.exists() and not args.overwrite and not args.compare_only:
        parser.error(f"结果文件已存在：{args.output}；可指定新文件名或使用 --overwrite")
    return args


def main() -> int:
    args = parse_args()
    started = perf_counter()
    try:
        source = load_source(args.input)
        print(
            f"读取 {len(source.orders)} 个生产订单、{source.repair_count} 条维修记录、{len(source.models)} 个机型",
            flush=True,
        )
        materials = None
        if not args.compare_only:
            query_started = perf_counter()
            materials = (
                None if args.offline else asyncio.run(fetch_sources(source, args))
            )
            if not args.offline:
                print(
                    f"接口查询耗时：{perf_counter() - query_started:.2f} 秒", flush=True
                )
            write_output(source, materials, args)
        comparison = compare_output(source, args, materials)
        write_comparison_report(comparison, source, args)
        for model in source.models:
            groups = [group for (m, _, _), group in source.groups.items() if m == model]
            production = total_production(source, model, args.quantity_field)
            print(
                f"{model}：维修{sum(g.count for g in groups)}条，汇总{len(groups)}行，整机数量{excel_number(production)}"
            )
        if materials is not None:
            print(f"接口发料记录：{materials.row_count}条")
            missing = sorted(
                {
                    f"{model}/{material}"
                    for model, _, material in source.groups
                    if material and (model, material) not in materials.quantities
                }
            )
            if missing:
                source.warnings.append("未匹配发料记录：" + "、".join(missing))
        for warning in source.warnings:
            print(f"提示：{warning}", file=sys.stderr)
        print(f"结果：{args.output}")
        print(f"对比报告：{args.report}")
        print(
            f"对比检查 {len(comparison.checks)} 项：失败 {comparison.failures} 项，"
            f"示例与原始数据冲突 {len(comparison.conflicts)} 项"
        )
        for conflict in comparison.conflicts:
            print(f"示例差异：{conflict}", file=sys.stderr)
        return comparison.exit_code
    except (ValueError, OSError, httpx.HTTPError) as exc:
        print(f"整合或对比失败：{exc}", file=sys.stderr)
        return 1
    finally:
        print(f"总耗时：{perf_counter() - started:.2f} 秒", flush=True)


if __name__ == "__main__":
    raise SystemExit(main())
