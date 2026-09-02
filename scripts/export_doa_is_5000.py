#!/usr/bin/env python3
"""Export DOA records marked as 5000 orders to an Excel workbook.

Run from the repository root with the temporary Excel dependency:

    uv run --project backend --with openpyxl python scripts/export_doa_is_5000.py

The script is read-only. It reads MongoDB settings from ``backend/.env`` (or
environment variables) and writes ``DOA_is_5000.xlsx`` to the directory in
which the command was started. Use ``--output`` to choose another path.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

from pymongo import MongoClient


PROJECT_ROOT = Path(__file__).resolve().parent.parent
BACKEND_DIR = PROJECT_ROOT / "backend"
DEFAULT_OUTPUT_NAME = "DOA_is_5000.xlsx"
COLLECTION_NAME = "scs_doa_results"
IS_5000_QUERY = {
    "$or": [
        {"is_5000": True},
        {"data.is_5000": True},
        {"scs_enrichment.is_5000": True},
    ]
}

sys.path.insert(0, str(BACKEND_DIR))

from app.core.config import get_settings  # noqa: E402


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="导出 is_5000=true 的 DOA 数据")
    parser.add_argument(
        "--output",
        type=Path,
        help=f"输出 Excel 路径，默认是当前目录/{DEFAULT_OUTPUT_NAME}",
    )
    return parser.parse_args()


def flatten(value: Any, prefix: str) -> dict[str, Any]:
    """Flatten nested mappings while keeping list values in one Excel cell."""
    if not isinstance(value, dict):
        return {prefix: value}

    result: dict[str, Any] = {}
    for key, item in value.items():
        child_prefix = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(item, dict):
            result.update(flatten(item, child_prefix))
        else:
            result[child_prefix] = item
    return result


def flatten_document(document: dict[str, Any]) -> dict[str, Any]:
    """Expose raw DOA fields and enrichment fields with unambiguous prefixes."""
    flattened: dict[str, Any] = {"_id": document.get("_id")}
    for key, value in document.items():
        if key in {"_id", "data", "scs_enrichment"}:
            continue
        flattened.update(flatten(value, f"document.{key}"))
    flattened.update(flatten(document.get("data") or {}, "data"))
    flattened.update(flatten(document.get("scs_enrichment") or {}, "scs_enrichment"))
    return flattened


def excel_value(value: Any) -> Any:
    """Convert MongoDB values to values accepted by openpyxl cells."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


def ordered_headers(keys: set[str]) -> list[str]:
    preferred = [
        "_id",
        "data.服务单号",
        "data.服务单号_id",
        "data.受理时间",
        "data.判定结果",
        "data.DOA类型",
        "data.处理状态",
        "data.问题批次数量",
        "scs_enrichment.sales_order",
        "scs_enrichment.material_sn",
        "scs_enrichment.accepted_date",
        "scs_enrichment.status",
        "scs_enrichment.is_5000",
        "data.is_5000",
    ]
    return [key for key in preferred if key in keys] + sorted(keys - set(preferred))


def write_excel(collection: Any, query: dict[str, Any], output_file: Path) -> int:
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise RuntimeError(
            "缺少 openpyxl，请使用命令：uv run --project backend --with openpyxl "
            "python scripts/export_doa_is_5000.py"
        ) from exc

    # Two cursor passes keep memory bounded while still producing a complete union of fields.
    keys: set[str] = {"_id"}
    for document in collection.find(query, no_cursor_timeout=True).sort("_id", 1):
        keys.update(flatten_document(document))
    headers = ordered_headers(keys)

    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("DOA_is_5000")
    worksheet.freeze_panes = "A2"

    header_fill = PatternFill("solid", fgColor="D9EAD3")
    header_font = Font(bold=True, color="1F2937")
    header_cells = []
    for header in headers:
        cell = WriteOnlyCell(worksheet, value=header)
        cell.fill = header_fill
        cell.font = header_font
        header_cells.append(cell)
    worksheet.append(header_cells)

    count = 0
    for document in collection.find(query, no_cursor_timeout=True).sort("_id", 1):
        row = flatten_document(document)
        worksheet.append([excel_value(row.get(header)) for header in headers])
        count += 1

    worksheet.auto_filter.ref = f"A1:{_excel_column(len(headers))}{count + 1}"

    output_file.parent.mkdir(parents=True, exist_ok=True)
    temporary_file = output_file.with_name(f".{output_file.name}.tmp")
    workbook.save(temporary_file)
    temporary_file.replace(output_file)
    return count


def _excel_column(number: int) -> str:
    """Return an Excel column name for a 1-based column number."""
    if number < 1:
        raise ValueError("Excel column number must be positive")
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def main() -> None:
    args = parse_args()
    output_file = (
        (args.output or Path.cwd() / DEFAULT_OUTPUT_NAME).expanduser().resolve()
    )

    # get_settings() resolves the repository's backend/.env relative to cwd.
    original_cwd = Path.cwd()
    os.chdir(BACKEND_DIR)
    try:
        settings = get_settings()
    finally:
        os.chdir(original_cwd)

    client = MongoClient(settings.MONGODB_URL, serverSelectionTimeoutMS=10_000)
    try:
        client.admin.command("ping")
        collection = client[settings.MONGODB_DB_NAME][COLLECTION_NAME]
        count = write_excel(collection, IS_5000_QUERY, output_file)
        print(f"已导出 {count} 条 DOA is_5000=true 记录：{output_file}")
    finally:
        client.close()


if __name__ == "__main__":
    main()
